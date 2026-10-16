"""
Modo planilha SEM o Microsoft Excel: lê arquivos .xlsx/.xlsm diretamente (openpyxl).

Usado quando o computador não tem o Excel (ou o pywin32). Oferece a mesma interface de
excel_com.PastaAberta / TabelaExcel, então o resto do programa não percebe a diferença.
O arquivo é aberto só para leitura: o programa nunca grava na planilha do usuário.
"""

import logging
from datetime import datetime
from pathlib import Path

from . import excel_com, planilha, utils

log = logging.getLogger("dare.planilha_arquivo")

EXTENSOES = (".xlsx", ".xlsm")


class PastaArquivo:
    """Pasta de trabalho lida do disco. Mesma interface de excel_com.PastaAberta."""

    def __init__(self, caminho: str | Path):
        try:
            import openpyxl
        except ImportError as exc:
            raise excel_com.ExcelIndisponivel(
                "Não há Excel neste computador e o leitor de planilhas (openpyxl) não está instalado."
            ) from exc
        arquivo = Path(str(caminho).strip().strip('"')).expanduser()
        if not arquivo.is_file():
            raise FileNotFoundError(f"Arquivo não encontrado: {arquivo}")
        if arquivo.suffix.lower() not in EXTENSOES:
            raise excel_com.ExcelIndisponivel(
                f"Sem o Excel instalado, só dá para ler .xlsx e .xlsm (este é {arquivo.suffix})."
            )
        try:
            # data_only: usa o último valor calculado das fórmulas (como o Excel mostra)
            self._wb = openpyxl.load_workbook(arquivo, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise excel_com.ExcelIndisponivel(
                f"Não consegui abrir '{arquivo.name}' (protegido por senha ou corrompido?): {exc}"
            ) from exc
        self.nome = arquivo.name
        self.caminho = str(arquivo.resolve())
        self.aberta_por_nos = True
        log.info("Planilha lida do disco (sem Excel): %s", arquivo)

    def abas(self) -> list[str]:
        return [ws.title for ws in self._wb.worksheets]

    def tabela(self, aba: str) -> planilha.TabelaMemoria:
        try:
            ws = self._wb[aba]
        except KeyError as exc:
            raise ValueError(f"A aba '{aba}' não existe nessa planilha.") from exc
        dados = [list(linha) for linha in ws.iter_rows(values_only=True)]
        ocultas = {n for n, dim in ws.row_dimensions.items() if dim.hidden}
        return planilha.TabelaMemoria(dados, nome=self.nome, aba=ws.title, ocultas=ocultas)

    def fechar(self) -> None:
        try:
            self._wb.close()
        except Exception:  # noqa: BLE001
            pass


def criar_modelo(caminho: str | Path) -> Path:
    """Cria a planilha modelo sem precisar do Excel (mesmo conteúdo do modelo do excel_com)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "DARE"
    ws.append(excel_com.MODELO_CABECALHOS)
    ws.append(["0012.123456/2026-12", "11222333000181", "52998224725", "11222333000181", "102026",
               datetime.combine(utils.vencimento_padrao(), datetime.min.time()), 100.0])
    for celula in ws[1]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="C4D8E6")
    formatos = ["@", "@", "@", "@", "@", "dd/mm/yyyy", "#,##0.00"]
    for col, formato in enumerate(formatos, 1):
        letra = planilha.numero_para_letra(col)
        ws.column_dimensions[letra].width = 22
        for linha in range(2, 1001):
            ws.cell(linha, col).number_format = formato
    ajuda = wb.create_sheet("Instruções")
    for i, texto in enumerate(excel_com.MODELO_INSTRUCOES, 1):
        ajuda.cell(i, 1, texto)
    ajuda.cell(1, 1).font = Font(bold=True)
    ajuda.column_dimensions["A"].width = 90
    wb.save(destino)
    log.info("Planilha modelo criada (sem Excel) em %s", destino)
    return destino


def excel_disponivel() -> bool:
    """True se o pywin32 está instalado (o Excel em si só é confirmado ao usar)."""
    try:
        import win32com.client  # noqa: F401
        return True
    except ImportError:
        return False


def abrir(caminho: str | Path):
    """
    Abre uma planilha pelo caminho: pelo Excel, se houver (lê .xls/.xlsb e fórmulas atualizadas);
    senão, direto do arquivo.
    """
    if excel_disponivel():
        try:
            return excel_com.abrir_por_caminho(caminho)
        except excel_com.ExcelIndisponivel:
            if Path(str(caminho)).suffix.lower() not in EXTENSOES:
                raise
            log.info("Excel indisponível; lendo o arquivo diretamente.")
    return PastaArquivo(caminho)


def criar_modelo_qualquer(caminho: str | Path) -> Path:
    """Cria o modelo pelo Excel se houver; senão, sem ele."""
    if excel_disponivel():
        try:
            return excel_com.criar_modelo(caminho)
        except excel_com.ExcelIndisponivel:
            log.info("Excel indisponível; criando o modelo sem ele.")
    return criar_modelo(caminho)
